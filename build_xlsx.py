from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule
from datetime import date
wb = Workbook()
INK="0B2239"; PAPER="F6F5F0"; TURQ="19B3D9"; TURQ_T="0C7893"; ROSE="FF4F8B"; SIGNAL="FFD233"; MIST="CDD5DD"; SLATE="4B5A6B"
thin = Side(style="thin", color=MIST)
hdr = NamedStyle(name="Clixite en-tête"); hdr.font = Font(name="Space Grotesk", bold=True, size=11, color=INK); hdr.fill = PatternFill("solid", fgColor=PAPER); hdr.alignment = Alignment(vertical="center", horizontal="left", indent=1); hdr.border = Border(bottom=Side(style="medium", color=INK))
body = NamedStyle(name="Clixite corps"); body.font = Font(name="Inter", size=10, color=INK); body.alignment = Alignment(vertical="center", indent=1); body.border = Border(bottom=thin)
num = NamedStyle(name="Clixite nombre"); num.font = Font(name="JetBrains Mono", size=10, color=INK); num.alignment = Alignment(horizontal="right", vertical="center", indent=1); num.border = Border(bottom=thin); num.number_format = '#,##0.00 "€"'
for s in (hdr, body, num): wb.add_named_style(s)
ws = wb.active; ws.title = "Suivi"; ws.sheet_view.showGridLines = False; ws.row_dimensions[1].height = 54
img = XLImage("charte/downloads/logo/clixite-relief-color-transparent-800px.png"); img.width = 190; img.height = int(190*257/675); ws.add_image(img, "B1")
ws["B3"] = "Suivi des actions"; ws["B3"].font = Font(name="Space Grotesk", bold=True, size=18, color=INK)
ws["B4"] = "ENTERPRISE & SOLUTION ARCHITECTURE  ·  [Client]  ·  [Date]"; ws["B4"].font = Font(name="Space Grotesk", size=9, color=TURQ_T, bold=True)
cols = ["Action", "Responsable", "Échéance", "Statut", "Budget", "Commentaire"]; widths = [42, 22, 14, 14, 14, 40]
for i, (c, w) in enumerate(zip(cols, widths), 2):
    cell = ws.cell(row=6, column=i, value=c); cell.style = "Clixite en-tête"; ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[6].height = 26
rows = [["Intégrer World-Check au flux de qualification", "Nicolas Simon", "2026-11-15", "En cours", 4800, "Sandbox obtenue"],
        ["Former l'équipe achats au light track", "Responsable achats", "2026-11-30", "À planifier", 2400, ""],
        ["Valider les playbooks avec le risque", "CRO", "2026-12-10", "Fait", 0, "Validé le 20/10"],
        ["Rédiger la policy fournisseurs critiques", "Nicolas Simon", "2026-12-20", "En retard", 7200, "Dépend du point 1"]]
for r, row in enumerate(rows, 7):
    for c, v in enumerate(row, 2):
        cell = ws.cell(row=r, column=c)
        if c == 4: cell.value = date.fromisoformat(v); cell.style = "Clixite corps"; cell.number_format = "DD/MM/YYYY"
        elif c == 6: cell.value = v; cell.style = "Clixite nombre"
        else: cell.value = v; cell.style = "Clixite corps"
    ws.row_dimensions[r].height = 22
last = 6 + len(rows)
ws.cell(row=last+1, column=5, value="Total").style = "Clixite en-tête"
t = ws.cell(row=last+1, column=6, value=f"=SUM(F7:F{last})"); t.style = "Clixite nombre"; t.font = Font(name="JetBrains Mono", size=10, bold=True, color=INK)
rng = f"E7:E{last}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fait"'], fill=PatternFill("solid", fgColor="D6F0F7"), font=Font(color=TURQ_T, bold=True)))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"En retard"'], fill=PatternFill("solid", fgColor="FFE1EC"), font=Font(color="D81B60", bold=True)))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"En cours"'], fill=PatternFill("solid", fgColor="FFF4C2"), font=Font(color=INK, bold=True)))
ws.freeze_panes = "B7"; ws.column_dimensions["A"].width = 3; ws.print_title_rows = "6:6"; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
k = wb.create_sheet("KPI"); k.sheet_view.showGridLines = False; k.column_dimensions["A"].width = 3
k["B2"] = "Indicateurs"; k["B2"].font = Font(name="Space Grotesk", bold=True, size=18, color=INK)
for i, (lbl, val, fmt, col) in enumerate([("Délai moyen de qualification (jours)", 9, "0", TURQ), ("Achats sous contrôle (€)", 800000000, '#,##0 "€"', ROSE), ("Fournisseurs critiques évalués", 1, "0%", TURQ_T)]):
    r = 4 + i*4
    k.cell(row=r, column=2, value=lbl).font = Font(name="Inter", size=10, color=SLATE)
    c = k.cell(row=r+1, column=2, value=val); c.font = Font(name="Space Grotesk", bold=True, size=28, color=col); c.number_format = fmt; k.row_dimensions[r+1].height = 40
k.column_dimensions["B"].width = 48
p = wb.create_sheet("Palette"); p.sheet_view.showGridLines = False
for col, w in (("A",3),("B",22),("C",12),("D",60),("E",6)): p.column_dimensions[col].width = w
p["B2"] = "Palette Clixite"; p["B2"].font = Font(name="Space Grotesk", bold=True, size=18, color=INK)
for i, (nm, hx, use) in enumerate([("Encre", INK, "Fond sombre, titres, texte"), ("Papier", PAPER, "Fond clair"), ("Turquoise bleu", TURQ, "Marque, symbole, surfaces"), ("Turquoise profond", TURQ_T, "Texte et liens sur fond clair"), ("Rose", ROSE, "Accent de signature"), ("Signal", SIGNAL, "Point clé, jamais en texte"), ("Brume", MIST, "Filets, bordures"), ("Ardoise", SLATE, "Texte secondaire")]):
    r = 4 + i
    p.cell(row=r, column=2, value=nm).font = Font(name="Inter", size=10, color=INK)
    p.cell(row=r, column=3, value="#" + hx).font = Font(name="JetBrains Mono", size=10, color=INK)
    p.cell(row=r, column=4, value=use).font = Font(name="Inter", size=10, color=SLATE)
    p.cell(row=r, column=5).fill = PatternFill("solid", fgColor=hx)
wb.save("charte/downloads/office/Clixite-tableau-de-suivi.xlsx"); print("xlsx ok")
